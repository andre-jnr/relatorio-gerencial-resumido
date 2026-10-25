"""Plano de contas: lê a planilha de classificação dos pagamentos.

Colunas usadas (primeira linha): PAGAMENTOS (nome da conta contábil), GRUPO e PLANO DE CONTAS.
A planilha é relida sozinha quando o arquivo muda.
"""
import os
import re
import threading
import unicodedata

from openpyxl import load_workbook

import config

_trava = threading.Lock()
_estado = {"mtime": None, "contas": {}}


def normalizar(texto):
    """Sem acentos, maiúsculas e sem espaços: 'Água  e Luz' -> 'AGUAELUZ'."""
    if not texto:
        return ""
    sem_acento = "".join(c for c in unicodedata.normalize("NFD", str(texto)) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", "", sem_acento).upper()


def _ler(caminho):
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        linhas = wb.worksheets[0].iter_rows(values_only=True)
        cab = [normalizar(c) for c in next(linhas, [])]
        try:
            i_conta = cab.index("PAGAMENTOS")
        except ValueError:
            raise ValueError("A planilha de classificação precisa da coluna PAGAMENTOS na primeira linha.")
        i_grupo = cab.index("GRUPO") if "GRUPO" in cab else None
        i_plano = cab.index("PLANODECONTAS") if "PLANODECONTAS" in cab else None

        texto = lambda linha, i: str(linha[i]).strip() if i is not None and i < len(linha) and linha[i] else ""
        contas = {}
        for linha in linhas:
            chave = normalizar(texto(linha, i_conta))
            if chave:
                contas[chave] = {
                    "grupo": texto(linha, i_grupo) or "SEM GRUPO",
                    "plano": texto(linha, i_plano) or "SEM PLANO DE CONTAS",
                }
        return contas
    finally:
        wb.close()


def contas():
    """Mapa conta normalizada -> {plano, grupo}. Vazio se a planilha não existir."""
    caminho = config.PLANILHA_CLASSIFICACAO
    if not os.path.exists(caminho):
        return {}
    mtime = os.path.getmtime(caminho)
    with _trava:
        if mtime != _estado["mtime"]:
            _estado["contas"] = _ler(caminho)
            _estado["mtime"] = mtime
        return _estado["contas"]
